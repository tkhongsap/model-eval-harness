"""Tests for the documents exact-match evaluation layer.

Synthetic sheets throughout, built by :func:`build_sheet`, so the whole scoring and rendering
path runs with no SharePoint and no sample workbook to go missing. The one exception is
``TestRealWorkbook``, which exercises the same path against the labelled run artifact under
``debug/`` and skips cleanly when it is absent.

What is worth pinning here is the handful of decisions that separate this module from its
sentiment counterpart: money folds to two decimals, dates shed Excel's midnight suffix, CER is
case-sensitive and GT-conditioned, coverage is GT-conditioned, and format compliance judges the
prediction alone and skips blanks. Each has a test that fails loudly if a later edit reverses
it.
"""

import io
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from src.local_model.documents.internal_exact_match import (
    BLOCK_CER,
    BLOCK_EXTRACTION,
    BLOCK_FORMAT,
    BLOCK_HEADER,
    BLOCK_LEGEND,
    BLOCK_OVERALL,
    CER_COLUMNS,
    COMPARE_AMOUNT,
    COMPARE_BUYER,
    COMPARE_DOC,
    COMPARE_GROUPS,
    COMPARE_VENDOR,
    FAMILY_CLASSIFICATION,
    FAMILY_EXTRACTION,
    FAMILY_QUALITY,
    FORMAT_CHECKS,
    INVERTED_RATIO_COLUMNS,
    KEY_COLUMN,
    MATCH_FALSE,
    MATCH_TRUE,
    NOT_APPLICABLE,
    SCORED_COLUMNS,
    SIMILARITY_THRESHOLDS,
    Config,
    _coverage,
    _is_branch_code,
    _is_count_column,
    _is_iso_date,
    _is_latin_script,
    _is_ratio_column,
    _is_thai_script,
    _is_valid_thai_tax_id,
    _match_column,
    _merge,
    _normalise_date,
    _normalise_money,
    _normalise_text,
    _page_key,
    _score_cer_column,
    _score_format,
    compare_rows,
    evaluate,
    resolve_sheet_names,
    write_comparison_sheets,
    write_dashboard,
)

GT_SHEET = "Doc - Groundtruth"
RESULT_SHEET = "Doc - Internal Model Result"

COMPARE_KEYS = (COMPARE_DOC, COMPARE_BUYER, COMPARE_VENDOR, COMPARE_AMOUNT)

# 12 digits + the mod-11 check digit they produce (weighted sum 352, 352 % 11 == 0,
# (11 - 0) % 10 == 1). Kept as a constant so every fixture is format-compliant by default.
VALID_TAX_ID = "1234567890121"
# Same 12 digits, wrong check digit.
INVALID_TAX_ID = "1234567890122"

THAI_NAME = "บริษัท ทดสอบ จำกัด"
THAI_ADDRESS = "เลขที่ 1 ถนนทดสอบ กรุงเทพมหานคร"

WORKBOOK = Path(__file__).resolve().parents[1] / "debug" / "model_comparison_2026-08-09_18-25-52_gt.xlsx"
READ_KWARGS = {"dtype": str, "header": 0, "keep_default_na": False, "na_values": [""]}


def build_sheet(rows=4, **columns):
    """A minimal sheet carrying every column evaluate() reads.

    Defaults agree on every column and satisfy every format check, so a test supplies only the
    column it is about and any disagreement or violation in the result is the one it introduced.
    Column headers carry spaces, so overrides are passed ``**{"Document type": [...]}``.
    """
    if columns:
        rows = len(next(iter(columns.values())))

    frame = {
        "No": list(range(1, rows + 1)),
        KEY_COLUMN: [f"doc_{i}_p1.pdf#1" for i in range(rows)],
        "File name": [f"doc_{i}.pdf" for i in range(rows)],
        "Document name": ["Tax Invoice"] * rows,
        "Document type": ["Receipt"] * rows,
        "Buyer name th": [THAI_NAME] * rows,
        "Buyer address th": [THAI_ADDRESS] * rows,
        "Buyer name eng": ["Acme Ltd"] * rows,
        "Buyer address eng": ["1 Test Road Bangkok"] * rows,
        "Buyer tax id": [VALID_TAX_ID] * rows,
        "Buyer branch code": ["00000"] * rows,
        "Buyer branch name": ["สำนักงานใหญ่"] * rows,
        "Vendor name th": [THAI_NAME] * rows,
        "Vendor address th": [THAI_ADDRESS] * rows,
        "Vendor name eng": ["Vendor Co Ltd"] * rows,
        "Vendor address eng": ["2 Sample Street Bangkok"] * rows,
        "Vendor tax id": [VALID_TAX_ID] * rows,
        "Vendor branch code": ["00000"] * rows,
        "Vendor branch name": ["สำนักงานใหญ่"] * rows,
        "Tax invoice number": ["INV-001"] * rows,
        "Tax invoice date": ["2026-01-05"] * rows,
        "Total amount": ["100.00"] * rows,
        "Vat amount": ["7.00"] * rows,
        "Net amount": ["107.00"] * rows,
        "Copy": ["False"] * rows,
        "Payee signature flag": ["True"] * rows,
        "Authorized receiver signature flag": ["True"] * rows,
        "Authorized signatory signature flag": ["False"] * rows,
        "Withholding tax": ["3.00"] * rows,
        "Invoice number": ["1"] * rows,
        "Invoice amount": ["100.00"] * rows,
        "Vat invoice": ["7.00"] * rows,
        "Stamp": ["False"] * rows,
    }
    frame.update(columns)
    return pd.DataFrame(frame)


def frame_row(frame, key_column, value):
    """The single row of a block whose key column holds ``value``."""
    matches = frame[frame[key_column] == value]
    assert len(matches) == 1, f"{value!r} matched {len(matches)} rows"
    return matches.iloc[0]


def starting_workbook(gt, result):
    """The two source sheets as a workbook, the shape run() downloads."""
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        gt.to_excel(writer, sheet_name=GT_SHEET, index=False)
        result.to_excel(writer, sheet_name=RESULT_SHEET, index=False)
    return buffer.getvalue()


class TestNormalise:
    def test_money_folds_to_two_decimals(self):
        # The Decimal("100.0") / Decimal("100.00") Excel round-trip must not read as an error.
        assert _normalise_money("100.0") == _normalise_money("100.00") == "100.00"

    def test_money_tolerates_thousands_separators(self):
        assert _normalise_money("1,234.5") == _normalise_money("1234.50") == "1234.50"

    def test_money_beyond_formatting_is_exact(self):
        # A one-satang difference is a real extraction error, not noise.
        assert _normalise_money("100.00") != _normalise_money("100.01")

    def test_unparsable_money_falls_back_to_text(self):
        assert _normalise_money("12..3") == _normalise_text("12..3")
        assert _normalise_money("about 100") == "about 100"

    def test_date_sheds_the_excel_midnight_suffix(self):
        assert _normalise_date("2026-01-05 00:00:00") == "2026-01-05"
        assert _normalise_date("2026-01-05T00:00:00") == "2026-01-05"
        assert _normalise_date("2026-01-05") == "2026-01-05"

    def test_a_real_time_of_day_is_not_stripped(self):
        # Only the midnight the round-trip appends is formatting; anything else is content.
        assert _normalise_date("2026-01-05 12:30:00") != "2026-01-05"

    def test_text_folds_case_and_internal_whitespace(self):
        assert _normalise_text("  ACME   Ltd ") == _normalise_text("acme ltd") == "acme ltd"

    def test_blank_and_missing_fold_to_empty(self):
        assert _normalise_text(None) == _normalise_text("") == _normalise_text("   ") == ""
        assert _normalise_money(None) == _normalise_money("") == ""


class TestMatchColumn:
    def test_scores_and_verdicts_come_from_one_pass(self):
        gt = build_sheet(**{"Document type": ["Receipt", "TaxInvoice", "Other", "Receipt"]})
        result = build_sheet(**{"Document type": ["Receipt", "Receipt", "Other", "TaxInvoice"]})
        merged, _, _ = _merge(gt, result)

        report, verdicts = _match_column(merged, "Document type")

        assert (report.n, report.tp, report.fp) == (4, 2, 2)
        # The verdicts must be the same comparison the counts came from, not a second one.
        assert verdicts == [MATCH_TRUE, MATCH_FALSE, MATCH_TRUE, MATCH_FALSE]
        assert verdicts.count(MATCH_TRUE) == report.tp

    def test_money_formatting_difference_is_a_match(self):
        gt = build_sheet(**{"Total amount": ["1,234.5", "100.0"]})
        result = build_sheet(**{"Total amount": ["1234.50", "100.00"]})
        merged, _, _ = _merge(gt, result)

        report, verdicts = _match_column(merged, "Total amount")

        assert report.tp == 2
        assert verdicts == [MATCH_TRUE, MATCH_TRUE]

    def test_unparsable_money_is_compared_as_text_and_logged(self, capsys):
        gt = build_sheet(**{"Total amount": ["12..3", "100.00"]})
        result = build_sheet(**{"Total amount": ["12..3", "100.00"]})
        merged, _, _ = _merge(gt, result)

        report, verdicts = _match_column(merged, "Total amount")

        # Identical garbage still matches -- deterministic, never a crash -- and is counted.
        assert verdicts == [MATCH_TRUE, MATCH_TRUE]
        # capsys rather than caplog: structlog writes through its own logger factory, not the
        # stdlib logging caplog hooks, so a caplog assertion here passes vacuously.
        captured = capsys.readouterr()
        record = captured.err + captured.out
        assert "money.unparsable" in record

    def test_date_midnight_suffix_is_a_match(self):
        gt = build_sheet(**{"Tax invoice date": ["2026-01-05 00:00:00", "2026-02-01"]})
        result = build_sheet(**{"Tax invoice date": ["2026-01-05", "2026-02-01"]})
        merged, _, _ = _merge(gt, result)

        report, _ = _match_column(merged, "Tax invoice date")

        assert report.tp == 2

    def test_two_blanks_count_as_a_match_and_are_logged(self, capsys):
        gt = build_sheet(**{"Withholding tax": ["", "", "3.00"]})
        result = build_sheet(**{"Withholding tax": ["", "", "3.00"]})
        merged, _, _ = _merge(gt, result)

        report, verdicts = _match_column(merged, "Withholding tax")

        # The documented consequence of excluding nothing. Pinned so it cannot change unnoticed,
        # and so the warning that makes it visible cannot be dropped as noise.
        assert report.accuracy == 1.0
        assert verdicts == [MATCH_TRUE, MATCH_TRUE, MATCH_TRUE]

        captured = capsys.readouterr()
        record = captured.err + captured.out
        assert "blank_both_sides" in record
        assert "Withholding tax" in record

    def test_a_blank_on_one_side_only_is_a_mismatch(self):
        gt = build_sheet(**{"Invoice number": ["1", "2"]})
        result = build_sheet(**{"Invoice number": ["", "2"]})
        merged, _, _ = _merge(gt, result)

        report, _ = _match_column(merged, "Invoice number")

        assert report.tp == 1
        assert report.fp == 1


class TestSimilarityMatch:
    def test_the_thresholds_cover_exactly_the_eight_party_fields(self):
        # The business rule: names at 90%, addresses at 80%, nothing else by similarity.
        assert {c for c in SIMILARITY_THRESHOLDS if "name" in c} == {
            "Buyer name th",
            "Buyer name eng",
            "Vendor name th",
            "Vendor name eng",
        }
        assert all(SIMILARITY_THRESHOLDS[c] == 0.90 for c in SIMILARITY_THRESHOLDS if "name" in c)
        assert all(
            SIMILARITY_THRESHOLDS[c] == 0.80 for c in SIMILARITY_THRESHOLDS if "address" in c
        )
        assert len(SIMILARITY_THRESHOLDS) == 8

    def test_a_name_passes_at_ninety_percent_and_fails_below(self):
        # 10 characters: one substitution is similarity 0.90 (a pass, on the boundary), two
        # substitutions 0.80 (a fail).
        gt = build_sheet(**{"Buyer name eng": ["0123456789", "0123456789"]})
        result = build_sheet(**{"Buyer name eng": ["012345678X", "01234567XX"]})
        merged, _, _ = _merge(gt, result)

        report, verdicts = _match_column(merged, "Buyer name eng")

        assert verdicts == [MATCH_TRUE, MATCH_FALSE]
        assert (report.tp, report.fp) == (1, 1)

    def test_an_address_passes_at_eighty_percent_and_fails_below(self):
        # 10 characters: two substitutions is similarity 0.80 (a pass, on the boundary), three
        # substitutions 0.70 (a fail).
        gt = build_sheet(**{"Buyer address eng": ["0123456789", "0123456789"]})
        result = build_sheet(**{"Buyer address eng": ["01234567XX", "0123456XXX"]})
        merged, _, _ = _merge(gt, result)

        _, verdicts = _match_column(merged, "Buyer address eng")

        assert verdicts == [MATCH_TRUE, MATCH_FALSE]

    def test_case_and_spacing_fold_before_similarity(self):
        # The fold runs first, so no similarity budget is spent on case or spacing.
        gt = build_sheet(**{"Vendor name eng": ["ACME LTD"]})
        result = build_sheet(**{"Vendor name eng": ["acme   ltd"]})
        merged, _, _ = _merge(gt, result)

        _, verdicts = _match_column(merged, "Vendor name eng")

        assert verdicts == [MATCH_TRUE]

    def test_blank_rules_are_unchanged_under_similarity(self):
        gt = build_sheet(**{"Vendor address th": [THAI_ADDRESS, "", ""]})
        result = build_sheet(**{"Vendor address th": ["", "", THAI_ADDRESS]})
        merged, _, _ = _merge(gt, result)

        _, verdicts = _match_column(merged, "Vendor address th")

        # Blank both sides is still a match; a blank against a filled cell still fails --
        # character_error_rate gives both for free (similarity 1.0 and 0.0 respectively).
        assert verdicts == [MATCH_FALSE, MATCH_TRUE, MATCH_FALSE]

    def test_exact_columns_do_not_inherit_the_similarity_rule(self):
        # One substitution in ten characters would pass the name threshold; on an identifier it
        # must stay a mismatch -- one wrong digit is a wrong number.
        gt = build_sheet(**{"Tax invoice number": ["0123456789"]})
        result = build_sheet(**{"Tax invoice number": ["012345678X"]})
        merged, _, _ = _merge(gt, result)

        _, verdicts = _match_column(merged, "Tax invoice number")

        assert verdicts == [MATCH_FALSE]

    def test_the_extraction_block_names_the_rule_per_row(self):
        blocks = evaluate(build_sheet(), build_sheet())
        block = blocks[BLOCK_EXTRACTION]

        assert frame_row(block, "Column", "Buyer name eng")["Match rule"] == "similarity >= 0.90"
        assert (
            frame_row(block, "Column", "Vendor address th")["Match rule"] == "similarity >= 0.80"
        )
        assert frame_row(block, "Column", "Document type")["Match rule"] == "exact"

    def test_similarity_verdicts_reconcile_with_the_dashboard(self):
        gt = build_sheet(**{"Buyer name eng": ["0123456789", "0123456789", "Acme Ltd"]})
        result = build_sheet(**{"Buyer name eng": ["012345678X", "01234XXXXX", "Acme Ltd"]})

        blocks = evaluate(gt, result)
        frame = compare_rows(gt, result)[COMPARE_BUYER]
        row = frame_row(blocks[BLOCK_EXTRACTION], "Column", "Buyer name eng")

        assert list(frame["Buyer name eng Compare"]).count(MATCH_TRUE) == row["TP"]
        assert list(frame["Buyer name eng Compare"]).count(MATCH_FALSE) == row["FP"]


class TestCoverage:
    def test_counts_are_conditioned_on_the_ground_truth(self):
        gt = build_sheet(**{"Invoice number": ["1", "2", "3", ""]})
        result = build_sheet(**{"Invoice number": ["9", "", "3", "4"]})
        merged, _, _ = _merge(gt, result)

        # Row 0 is covered despite being WRONG -- coverage is completeness, not correctness --
        # and row 3's hallucinated value is not counted, because the human left it blank.
        assert _coverage(merged, "Invoice number") == (3, 2)

    def test_a_column_the_human_never_filled_has_nothing_to_cover(self):
        gt = build_sheet(**{"Withholding tax": ["", "", ""]})
        result = build_sheet(**{"Withholding tax": ["3.00", "", ""]})
        merged, _, _ = _merge(gt, result)

        assert _coverage(merged, "Withholding tax") == (0, 0)


class TestCer:
    def test_identical_text_scores_zero(self):
        gt = build_sheet(rows=3)
        merged, _, _ = _merge(gt, gt.copy())

        report, excluded = _score_cer_column(merged, "Buyer name th", 0.20)

        assert (report.mean, report.n, excluded) == (0.0, 3, 0)

    def test_blank_ground_truth_rows_are_excluded_and_counted(self, capsys):
        gt = build_sheet(**{"Buyer address eng": ["", "1 Test Road", ""]})
        result = build_sheet(**{"Buyer address eng": ["whatever", "1 Test Road", ""]})
        merged, _, _ = _merge(gt, result)

        report, excluded = _score_cer_column(merged, "Buyer address eng", 0.20)

        # Nothing to err against on a blank reference; the row measures nothing.
        assert (report.n, excluded) == (1, 2)
        assert report.mean == 0.0

        captured = capsys.readouterr()
        record = captured.err + captured.out
        assert "blank_ground_truth" in record

    def test_blank_prediction_against_filled_ground_truth_scores_one(self):
        gt = build_sheet(**{"Vendor name eng": ["Vendor Co"]})
        result = build_sheet(**{"Vendor name eng": [""]})
        merged, _, _ = _merge(gt, result)

        report, excluded = _score_cer_column(merged, "Vendor name eng", 0.20)

        # Every reference character is a deletion -- a missing extraction is a bad extraction,
        # not an unmeasured one.
        assert (report.mean, report.n, excluded) == (1.0, 1, 0)

    def test_cer_is_case_sensitive(self):
        # The exact-match column beside it already reports the case-blind view; CER exists to
        # count real character errors.
        gt = build_sheet(**{"Buyer name eng": ["Acme"]})
        result = build_sheet(**{"Buyer name eng": ["ACME"]})
        merged, _, _ = _merge(gt, result)

        report, _ = _score_cer_column(merged, "Buyer name eng", 0.20)

        assert report.mean > 0.0


class TestFormatChecks:
    def test_valid_mod11_tax_id_passes(self):
        assert _is_valid_thai_tax_id(VALID_TAX_ID)

    @pytest.mark.parametrize(
        "value",
        [
            INVALID_TAX_ID,  # right length, wrong check digit
            "123456789012",  # 12 digits
            "12345678901234",  # 14 digits
            "123456789012a",  # non-digit
        ],
    )
    def test_malformed_tax_ids_fail(self, value):
        assert not _is_valid_thai_tax_id(value)

    @pytest.mark.parametrize("value", ["2026-01-05", "2026-01-05 00:00:00", "2026-12-31"])
    def test_iso_dates_pass(self, value):
        assert _is_iso_date(value)

    @pytest.mark.parametrize(
        "value",
        [
            "05/01/2026",  # not ISO
            "2569-01-05",  # Buddhist-era year the model failed to convert
            "0000-01-01",  # the placeholder the prompt bans
            "2026-13-01",  # not a date
            "2026-01-05 12:30:00",  # a real time of day is not the round-trip suffix
        ],
    )
    def test_non_iso_dates_fail(self, value):
        assert not _is_iso_date(value)

    def test_script_purity(self):
        assert _is_latin_script("Acme Ltd 123, Bangkok")
        assert not _is_latin_script(f"Acme {THAI_NAME}")
        assert _is_thai_script(THAI_ADDRESS)
        assert not _is_thai_script("Acme Ltd")
        # A th field holding pure ASCII digits has no Thai script to show.
        assert not _is_thai_script("12345")
        assert not _is_thai_script(f"{THAI_NAME} Ltd")

    @pytest.mark.parametrize("value,expected", [("00000", True), ("12345", True), ("123", False), ("1234a", False)])
    def test_branch_codes(self, value, expected):
        assert _is_branch_code(value) is expected

    def test_blank_predictions_are_skipped_not_failed(self):
        gt = build_sheet(rows=4)
        result = build_sheet(**{"Buyer tax id": ["", VALID_TAX_ID, INVALID_TAX_ID, ""]})
        merged, _, _ = _merge(gt, result)

        block = _score_format(merged)
        row = block[(block["Column"] == "Buyer tax id")].iloc[0]

        # Completeness is Coverage's job; format judges only what was written.
        assert (row["Checked"], row["Passed"], row["Failed"]) == (2, 1, 1)
        assert row["Pass rate"] == 0.5

    def test_a_check_nothing_reached_reads_not_applicable(self):
        gt = build_sheet(rows=2)
        result = build_sheet(**{"Tax invoice date": ["", ""]})
        merged, _, _ = _merge(gt, result)

        block = _score_format(merged)
        row = frame_row(block, "Column", "Tax invoice date")

        assert row["Checked"] == 0
        assert row["Pass rate"] == NOT_APPLICABLE

    def test_registry_covers_thirteen_checks(self):
        assert len(FORMAT_CHECKS) == 13


class TestEvaluateBlocks:
    def test_blocks_arrive_in_sheet_order(self):
        blocks = evaluate(build_sheet(), build_sheet())

        assert list(blocks) == [
            BLOCK_HEADER,
            BLOCK_OVERALL,
            BLOCK_EXTRACTION,
            BLOCK_CER,
            BLOCK_FORMAT,
            BLOCK_LEGEND,
        ]

    def test_extraction_block_carries_every_scored_column_in_order(self):
        blocks = evaluate(build_sheet(), build_sheet())

        assert list(blocks[BLOCK_EXTRACTION]["Column"]) == list(SCORED_COLUMNS)
        assert list(blocks[BLOCK_EXTRACTION].columns) == [
            "Column",
            "N",
            "TP",
            "FP",
            "FN",
            "TN",
            "Accuracy",
            "Precision",
            "Recall",
            "F1",
            "GT filled",
            "Coverage",
            "Match rule",
        ]

    def test_cer_block_carries_the_eight_party_columns(self):
        blocks = evaluate(build_sheet(), build_sheet())

        assert list(blocks[BLOCK_CER]["Column"]) == list(CER_COLUMNS)

    def test_self_comparison_scores_one_everywhere(self):
        """The strongest check available: a sheet against itself must be perfect."""
        sheet = build_sheet(
            **{
                "Document type": ["Receipt", "TaxInvoice", "Other", "Receipt"],
                "Total amount": ["1,234.5", "99.9", "0.00", "100.00"],
            }
        )

        blocks = evaluate(sheet, sheet.copy())

        block = blocks[BLOCK_EXTRACTION]
        assert list(block["TP"]) == list(block["N"])
        assert set(block["FP"]) == {0}
        for column in ("Accuracy", "Precision", "Recall", "F1", "Coverage"):
            assert set(block[column]) == {1.0}, column

        assert set(blocks[BLOCK_CER]["Mean CER"]) == {0.0}

        overall = blocks[BLOCK_OVERALL]
        for family in (FAMILY_CLASSIFICATION, FAMILY_EXTRACTION):
            row = frame_row(overall, "Family", family)
            assert row["Accuracy"] == 1.0
            assert row["F1"] == 1.0

    def test_n_is_the_matched_row_count_on_every_column(self):
        gt = build_sheet(**{"Document type": ["Receipt", "not a type", "Other", ""]})
        result = build_sheet()

        blocks = evaluate(gt, result)

        # Nothing is excluded here -- the out-of-vocabulary and blank rows are mismatches, and
        # no column can quietly score itself on a smaller sample than the one beside it. The
        # confusion-matrix module makes the opposite choice, on purpose.
        assert set(blocks[BLOCK_EXTRACTION]["N"]) == {4}

    def test_quality_family_row_stays_not_applicable(self):
        blocks = evaluate(build_sheet(), build_sheet())
        row = frame_row(blocks[BLOCK_OVERALL], "Family", FAMILY_QUALITY)

        assert row["Accuracy"] == NOT_APPLICABLE
        assert row["F1"] == NOT_APPLICABLE
        assert "mean CER" in row["Note"]
        assert "format pass" in row["Note"]

    def test_coverage_of_an_unfilled_column_reads_not_applicable(self):
        gt = build_sheet(**{"Withholding tax": ["", "", "", ""]})
        result = build_sheet(**{"Withholding tax": ["", "", "", ""]})

        blocks = evaluate(gt, result)
        row = frame_row(blocks[BLOCK_EXTRACTION], "Column", "Withholding tax")

        # Blank==blank keeps Accuracy at 1.0 -- which is exactly why Coverage must say n/a
        # rather than pretending the column was extracted.
        assert row["Accuracy"] == 1.0
        assert row["GT filled"] == 0
        assert row["Coverage"] == NOT_APPLICABLE

    def test_legend_explains_the_documents_specific_metrics(self):
        blocks = evaluate(build_sheet(), build_sheet())
        legend = dict(
            zip(blocks[BLOCK_LEGEND]["Term"], blocks[BLOCK_LEGEND]["Meaning"], strict=True)
        )

        assert "right or wrong" in legend["Coverage"]
        assert "LOWER IS BETTER" in legend["CER (Character Error Rate)"]
        assert "PREDICTION ALONE" in legend["Format compliance"]
        assert "SAME NUMBER AS ACCURACY" in legend["Precision"]
        # The similarity rule changes what Accuracy means on eight rows; the legend must say
        # both the rule and its thresholds.
        assert "0.90" in legend["Match rule"]
        assert "0.80" in legend["Match rule"]


class TestCompareRows:
    def test_the_four_groups_cover_every_scored_column_exactly_once(self):
        spread = [column for columns in COMPARE_GROUPS.values() for column in columns]

        assert sorted(spread) == sorted(SCORED_COLUMNS)

    def test_one_row_per_matched_key_in_join_order(self):
        frames = compare_rows(build_sheet(rows=3), build_sheet(rows=3))

        for key in COMPARE_KEYS:
            assert len(frames[key]) == 3
            assert list(frames[key]["No"]) == [1, 2, 3]

    def test_each_sheet_has_a_triplet_per_group_column(self):
        frames = compare_rows(build_sheet(), build_sheet())

        for key, columns in COMPARE_GROUPS.items():
            frame_columns = list(frames[key].columns)
            assert frame_columns[:2] == ["No", KEY_COLUMN]
            assert len(frame_columns) == 2 + len(columns) * 3

    def test_cells_hold_the_raw_value_not_the_folded_one(self):
        gt = build_sheet(**{"Total amount": ["100.0", "1,234.5"]})
        result = build_sheet(**{"Total amount": ["100.00", "1234.50"]})

        frame = compare_rows(gt, result)[COMPARE_AMOUNT]

        # The sheet is evidence, so it must show what the workbook held: the verdict says T
        # while the cells print the unfolded pair.
        assert list(frame["Total amount GT"]) == ["100.0", "1,234.5"]
        assert list(frame["Total amount AI"]) == ["100.00", "1234.50"]
        assert list(frame["Total amount Compare"]) == [MATCH_TRUE, MATCH_TRUE]

    def test_verdicts_reconcile_with_the_dashboard(self):
        gt = build_sheet(**{"Document type": ["Receipt", "TaxInvoice", "Other", "Receipt"]})
        result = build_sheet(**{"Document type": ["Receipt", "Receipt", "Other", "TaxInvoice"]})

        blocks = evaluate(gt, result)
        frame = compare_rows(gt, result)[COMPARE_DOC]
        row = frame_row(blocks[BLOCK_EXTRACTION], "Column", "Document type")

        # The compare sheet is the evidence for the block's counts; if they can disagree, it is
        # evidence for something else.
        assert list(frame["Document type Compare"]).count(MATCH_TRUE) == row["TP"]
        assert list(frame["Document type Compare"]).count(MATCH_FALSE) == row["FP"]

    def test_empty_join_keeps_the_column_shape(self):
        gt = build_sheet(rows=2)
        result = build_sheet(rows=2)
        result[KEY_COLUMN] = ["other_a.pdf#1", "other_b.pdf#1"]

        frames = compare_rows(gt, result)

        # Without this the sheet would lose its header and read as a successful empty run.
        for key, columns in COMPARE_GROUPS.items():
            assert len(frames[key]) == 0
            assert len(frames[key].columns) == 2 + len(columns) * 3

    def test_the_key_column_prints_the_ground_truth_spelling(self):
        gt = build_sheet(rows=2)
        result = build_sheet(rows=2)
        result[KEY_COLUMN] = [key.replace(".pdf#", ".png#") for key in result[KEY_COLUMN]]

        frame = compare_rows(gt, result)[COMPARE_DOC]

        # The two sides now spell the page differently and the join folds that away, so the
        # sheet has to pick one: the ground truth's, which is what the reader looks the row
        # up by and what the banner's "N of M ground-truth rows" counts.
        assert list(frame[KEY_COLUMN]) == ["doc_0_p1.pdf#1", "doc_1_p1.pdf#1"]


class TestPageKey:
    @pytest.mark.parametrize(
        "key",
        ["doc_p1.pdf#1", "doc_p1.png#1", "doc_p1.jpg#1", "doc_p1.jpeg#1", "doc_p1.PNG#1"],
    )
    def test_every_page_extension_folds_to_one_identity(self, key):
        # The whole point: the internal split renders a PDF page to PNG where the ground
        # truth names the same page PDF, and the join saw two unrelated rows.
        assert _page_key(key) == "doc_p1#1"

    def test_the_line_position_is_matched_exactly(self):
        # Only the extension folds. A page where the model read fewer printed lines than the
        # human recorded must still surface as an unmatched key, not as a silent alignment.
        assert _page_key("doc_p2.png#3") == "doc_p2#3"
        assert _page_key("doc_p2.png#1") != _page_key("doc_p2.png#2")

    def test_a_key_with_no_position_keeps_its_shape(self):
        assert _page_key("doc_p1.png") == "doc_p1"

    def test_a_stem_containing_a_dot_keeps_it(self):
        # Path().stem -- the sentiment modules' fold -- returns "INV_2026_p1" for both of
        # these, which is why the fold is a whitelist of real page extensions instead.
        assert _page_key("INV_2026.03_p1.png#2") == "INV_2026.03_p1#2"
        assert _page_key("INV_2026.03_p1#2") == "INV_2026.03_p1#2"

    def test_a_path_separator_survives(self):
        # Path().stem would drop the directory and merge two distinct pages into one key.
        assert _page_key("a/doc_p1.png#1") == "a/doc_p1#1"

    def test_blanks_fold_to_the_empty_string(self):
        assert _page_key(None) == ""
        assert _page_key(float("nan")) == ""
        assert _page_key("   ") == ""

    def test_the_fold_is_idempotent(self):
        # A ground truth already keyed without an extension has to join too.
        once = _page_key("doc_p1.png#1")
        assert _page_key(once) == once


class TestMerge:
    def test_duplicate_keys_multiply_rows_and_are_logged(self, capsys):
        gt = build_sheet(rows=3)
        gt.loc[1, KEY_COLUMN] = gt.loc[0, KEY_COLUMN]
        result = build_sheet(rows=3)

        merged, _, _ = _merge(gt, result)

        # The cartesian join is kept -- a duplicate is an upstream bug, and dropping it here
        # would hide the very row count that reveals it.
        assert len(merged) == 3
        captured = capsys.readouterr()
        record = captured.err + captured.out
        assert "duplicate_keys" in record

    def test_unmatched_keys_are_counted_per_side(self):
        gt = build_sheet(rows=3)
        result = build_sheet(rows=2)

        merged, gt_only, result_only = _merge(gt, result)

        assert (len(merged), gt_only, result_only) == (2, 1, 0)

    def test_a_png_result_joins_a_pdf_ground_truth(self):
        # The regression. The internal run renders each PDF page to a .png while the ground
        # truth names it .pdf, so this join used to come back empty and the dashboard printed
        # N=0 on every row without raising.
        gt = build_sheet(rows=3)
        result = build_sheet(rows=3)
        result[KEY_COLUMN] = [key.replace(".pdf#", ".png#") for key in result[KEY_COLUMN]]

        merged, gt_only, result_only = _merge(gt, result)

        assert (len(merged), gt_only, result_only) == (3, 0, 0)

    def test_page_names_differing_only_in_case_do_not_join(self):
        # Deliberate: the fold drops the extension and nothing else. Two source files
        # differing in case are two source files.
        gt = build_sheet(rows=2)
        result = build_sheet(rows=2)
        result[KEY_COLUMN] = [key.replace("doc_", "Doc_") for key in gt[KEY_COLUMN]]

        merged, _, _ = _merge(gt, result)

        assert len(merged) == 0

    def test_unmatched_keys_are_logged_with_a_sample(self, capsys):
        gt = build_sheet(rows=3)
        result = build_sheet(rows=2)

        _merge(gt, result)

        captured = capsys.readouterr()
        record = captured.err + captured.out
        assert "join.unmatched" in record
        # The keys themselves, not just the counts: a count cannot tell a key-format mismatch
        # from a page where the model read fewer printed lines than the human recorded.
        assert "doc_2_p1#1" in record

    def test_an_empty_join_is_logged_at_error_and_returned(self, capsys):
        gt = build_sheet(rows=2)
        result = build_sheet(rows=2)
        result[KEY_COLUMN] = ["other_a.pdf#1", "other_b.pdf#1"]

        merged, _, _ = _merge(gt, result)

        # Returned rather than raised -- the caller still writes and uploads its dashboard --
        # so the log line is the only thing standing between a sheet of zeros and a silent one.
        assert len(merged) == 0
        captured = capsys.readouterr()
        record = captured.err + captured.out
        assert "join.empty" in record


class TestColumnFormatting:
    @pytest.mark.parametrize("name", ["Accuracy", "Coverage", "Pass rate", "F1"])
    def test_the_higher_is_better_ratios_are_ratios(self, name):
        assert _is_ratio_column(name)
        assert not _is_count_column(name)

    def test_the_cer_pass_rate_header_matches_by_prefix(self):
        # The header carries its threshold, so an exact-name set cannot match it.
        assert _is_ratio_column("CER pass rate (<= 0.20)")

    @pytest.mark.parametrize("name", ["N", "GT filled", "Checked", "Passed", "Failed", "TP"])
    def test_the_counts_are_counts(self, name):
        assert _is_count_column(name)
        assert not _is_ratio_column(name)

    def test_raw_cer_columns_take_the_inverted_scale(self):
        # 0.04 is an excellent extraction; on the shared scale it would paint red.
        assert "Mean CER" in INVERTED_RATIO_COLUMNS
        assert "Worst CER (highest)" in INVERTED_RATIO_COLUMNS
        assert not _is_ratio_column("Mean CER")


BASES = (
    "Evaluation Dashboard",
    "Doc Compare Result",
    "Buyer Compare Result",
    "Vendor Compare Result",
    "Amount Compare Result",
)


class TestResolveSheetNames:
    def test_returns_the_bases_when_all_are_free(self):
        assert resolve_sheet_names(["Sheet1"], BASES) == list(BASES)

    def test_one_taken_name_moves_all_five(self):
        resolved = resolve_sheet_names(["Buyer Compare Result"], BASES)

        # The five sheets are only readable as a set. A dashboard paired with a verdict sheet
        # from a different run is worse than an ugly suffix.
        assert resolved == [f"{base}_1" for base in BASES]

    def test_config_supplies_all_five_names(self):
        config = Config()

        assert [
            config.eval_sheet_name,
            config.doc_compare_sheet_name,
            config.buyer_compare_sheet_name,
            config.vendor_compare_sheet_name,
            config.amount_compare_sheet_name,
        ] == list(BASES)


@pytest.fixture(scope="module")
def rendered():
    """One workbook carrying all five appended sheets, built once for the rendering tests."""
    gt = build_sheet(
        **{"Document name": ["Tax Invoice", "Receipt Copy", "Tax Invoice", "Tax Invoice"]}
    )
    result = build_sheet()
    original = starting_workbook(gt, result)

    blocks = evaluate(gt, result)
    frames = compare_rows(gt, result)
    names = resolve_sheet_names(load_workbook(io.BytesIO(original)).sheetnames, BASES)
    updated = write_dashboard(original, blocks, names[0])
    updated = write_comparison_sheets(
        updated, frames, dict(zip(COMPARE_KEYS, names[1:], strict=True))
    )
    return original, updated, names


class TestRenderedWorkbook:
    def test_all_five_sheets_are_appended(self, rendered):
        _, updated, names = rendered

        book = load_workbook(io.BytesIO(updated))

        assert book.sheetnames == [GT_SHEET, RESULT_SHEET, *names]

    def test_source_sheets_survive_untouched(self, rendered):
        original, updated, _ = rendered

        before = pd.read_excel(io.BytesIO(original), sheet_name=GT_SHEET)
        after = pd.read_excel(io.BytesIO(updated), sheet_name=GT_SHEET)

        assert before.equals(after)

    def test_compare_sheet_has_a_two_row_merged_header(self, rendered):
        _, updated, names = rendered
        sheet = load_workbook(io.BytesIO(updated))[names[1]]

        assert sheet.cell(row=1, column=1).value == "No"
        assert sheet.cell(row=1, column=2).value == KEY_COLUMN
        assert sheet.cell(row=1, column=3).value == COMPARE_GROUPS[COMPARE_DOC][0]
        # Row 1 names the group, row 2 names the three cells under it.
        assert [sheet.cell(row=2, column=c).value for c in (3, 4, 5)] == ["GT", "AI", "Compare"]
        assert "C1:E1" in {str(r) for r in sheet.merged_cells.ranges}
        assert "A1:A2" in {str(r) for r in sheet.merged_cells.ranges}

    def test_mismatch_cells_are_filled_red_and_matches_green(self, rendered):
        _, updated, names = rendered
        sheet = load_workbook(io.BytesIO(updated))[names[1]]

        # Column 5 is "Document name Compare"; the fixture mismatches row 2 only.
        verdicts = {sheet.cell(row=r, column=5).value: r for r in range(3, 7)}
        matched = sheet.cell(row=verdicts[MATCH_TRUE], column=5)
        missed = sheet.cell(row=verdicts[MATCH_FALSE], column=5)

        assert missed.fill.fgColor.rgb.endswith("FFC7CE")
        assert matched.fill.fgColor.rgb.endswith("C6EFCE")

    def test_dashboard_formats_n_as_an_integer_and_accuracy_at_four_places(self, rendered):
        _, updated, names = rendered
        sheet = load_workbook(io.BytesIO(updated))[names[0]]

        header = next(row for row in sheet.iter_rows() if row[0].value == "Column")
        first_data = header[0].row + 1

        assert [cell.value for cell in header[:6]] == ["Column", "N", "TP", "FP", "FN", "TN"]
        assert sheet.cell(row=first_data, column=2).number_format == "0"
        assert sheet.cell(row=first_data, column=7).number_format == "0.0000"

    def test_cer_columns_carry_an_inverted_colour_scale(self, rendered):
        _, updated, names = rendered
        sheet = load_workbook(io.BytesIO(updated))[names[0]]

        starts = set()
        for rules in sheet.conditional_formatting:
            for rule in rules.rules:
                if rule.colorScale is not None:
                    starts.add(str(rule.colorScale.color[0].rgb))

        # The normal scale starts red (F8696B); only the raw CER columns start green (63BE7B).
        # Both must be present, or either the scales or the inversion went missing.
        assert any(start.endswith("F8696B") for start in starts)
        assert any(start.endswith("63BE7B") for start in starts)

    def test_writing_twice_produces_a_second_suffixed_set(self, rendered):
        _, updated, _ = rendered
        gt = build_sheet(rows=2)
        result = build_sheet(rows=2)

        second = resolve_sheet_names(load_workbook(io.BytesIO(updated)).sheetnames, BASES)
        again = write_dashboard(updated, evaluate(gt, result), second[0])
        again = write_comparison_sheets(
            again, compare_rows(gt, result), dict(zip(COMPARE_KEYS, second[1:], strict=True))
        )

        assert second == [f"{base}_1" for base in BASES]
        assert load_workbook(io.BytesIO(again)).sheetnames[-5:] == second

    def test_writing_over_an_existing_compare_sheet_raises(self, rendered):
        _, updated, names = rendered
        gt = build_sheet(rows=2)

        with pytest.raises(ValueError, match="already exists"):
            write_comparison_sheets(
                updated,
                compare_rows(gt, gt.copy()),
                dict(zip(COMPARE_KEYS, names[1:], strict=True)),
            )


@pytest.fixture(scope="module")
def real_sheets():
    """The labelled run artifact's two sheets, when the file is present."""
    if not WORKBOOK.exists():
        pytest.skip(f"sample workbook not present: {WORKBOOK}")
    book = pd.ExcelFile(WORKBOOK, engine="openpyxl")
    if RESULT_SHEET not in book.sheet_names:
        pytest.skip(f"sheet not present in sample workbook: {RESULT_SHEET}")
    gt = book.parse(GT_SHEET, **READ_KWARGS)
    result = book.parse(RESULT_SHEET, **READ_KWARGS)
    return gt, result


class TestRealWorkbook:
    def test_the_join_is_not_empty(self, real_sheets):
        gt, result = real_sheets
        merged, _, _ = _merge(gt, result)

        assert len(merged) > 0

    def test_evaluate_runs_offline_over_the_real_frames(self, real_sheets):
        gt, result = real_sheets
        blocks = evaluate(gt, result)

        assert len(blocks[BLOCK_EXTRACTION]) == len(SCORED_COLUMNS)
        assert len(blocks[BLOCK_CER]) == len(CER_COLUMNS)
        assert len(blocks[BLOCK_FORMAT]) == len(FORMAT_CHECKS)
        # Every row matched: the ground truth was built from this run's own pages.
        row = frame_row(blocks[BLOCK_EXTRACTION], "Column", "Document type")
        assert row["N"] == len(_merge(gt, result)[0])

    def test_dashboard_round_trips_onto_the_real_workbook(self, real_sheets):
        gt, result = real_sheets
        original = starting_workbook(gt, result)

        blocks = evaluate(gt, result)
        updated = write_dashboard(original, blocks, "Evaluation Dashboard")

        assert "Evaluation Dashboard" in load_workbook(io.BytesIO(updated)).sheetnames
