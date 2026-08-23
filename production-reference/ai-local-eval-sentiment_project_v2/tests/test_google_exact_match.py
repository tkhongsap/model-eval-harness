"""Tests for the exact-match evaluation layer.

Synthetic sheets throughout, built by :func:`build_sheet`, so the whole scoring and rendering path
runs with no SharePoint, no Vertex and no sample workbook to go missing.

What is worth pinning here is the handful of decisions that separate this module from
``google_confusion_matrix``: the comparison folds case, nothing is ever excluded, and the compare
sheets must agree with the dashboard row for row. Each has a test that fails loudly if a later
edit reverses it.
"""

import io

import pandas as pd
import pytest
from openpyxl import load_workbook

from src.google_model.sentiment.google_exact_match import (
    BLOCK_CALL_TYPE_SUMMARY,
    BLOCK_LEGEND,
    BLOCK_OVERALL,
    BLOCK_QA,
    BLOCK_SENTIMENT,
    COMPARE_CALL_TYPE,
    COMPARE_QA,
    COMPARE_SENTIMENT,
    FAMILY_CALL_TYPE,
    FAMILY_QA,
    FAMILY_SENTIMENT,
    KEY_COLUMN,
    MATCH_FALSE,
    MATCH_TRUE,
    NOT_APPLICABLE,
    QA_CRITERIA,
    SENTIMENT_COLUMNS,
    Config,
    _is_count_column,
    _is_ratio_column,
    _match_column,
    _merge,
    _normalise,
    compare_rows,
    evaluate,
    resolve_sheet_names,
    write_comparison_sheets,
    write_dashboard,
)

GT_SHEET = "Voice - Groundtruth"
RESULT_SHEET = "Voice - Google Result"

COMPARE_KEYS = (COMPARE_QA, COMPARE_SENTIMENT, COMPARE_CALL_TYPE)


# Deterministic and dimension-stable: every text embeds to the same vector, so every cosine is
# exactly 1.0. Summary scoring is unchanged from google_confusion_matrix and has its own tests.
def stub_embedder(texts):
    return [[1.0, 0.0, 0.0] for _ in texts]


def build_sheet(rows=4, **columns):
    """A minimal sheet carrying every column evaluate() reads.

    Defaults agree on every column, so a test supplies only the column it is about and any
    disagreement in the result is the one it introduced.
    """
    if columns:
        rows = len(next(iter(columns.values())))

    frame = {KEY_COLUMN: [f"call_{i}.wav" for i in range(rows)]}
    frame.update({column: ["Meet"] * rows for column in QA_CRITERIA})
    frame.update({column: ["Neutral"] * rows for column in SENTIMENT_COLUMNS})
    frame["call_type"] = ["Enquiry"] * rows
    frame["summary_story"] = ["a summary"] * rows
    frame.update(columns)
    return pd.DataFrame(frame)


def frame_row(frame, key_column, value):
    """The single row of a block whose key column holds ``value``."""
    matches = frame[frame[key_column] == value]
    assert len(matches) == 1, f"{value!r} matched {len(matches)} rows"
    return matches.iloc[0]


def starting_workbook(gt, result):
    """The two source sheets as a workbook, the shape run() downloads."""
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        gt.to_excel(writer, sheet_name=GT_SHEET, index=False)
        result.to_excel(writer, sheet_name=RESULT_SHEET, index=False)
    return buffer.getvalue()


class TestNormalise:
    def test_case_and_surrounding_space_are_ignored(self):
        # The rule the reference workbook uses -- row 3 of "SR-RE Human vs AI R119" marks 'meet'
        # against 'Meet' as a match.
        assert _normalise(" meet ") == _normalise("Meet") == "meet"

    def test_blank_and_missing_both_fold_to_empty(self):
        assert _normalise(None) == _normalise("") == _normalise("   ") == ""

    def test_interior_difference_is_preserved(self):
        # Only the ends are stripped. A grade is one word, so an interior difference is a
        # different answer, not formatting.
        assert _normalise("not meet") != _normalise("notmeet")


class TestMatchColumn:
    def test_scores_and_verdicts_come_from_one_pass(self):
        gt = build_sheet(greeting_standard=["Meet", "Below", "N/A", "Meet"])
        result = build_sheet(greeting_standard=["Meet", "Meet", "N/A", "Below"])
        merged, _, _ = _merge(gt, result)

        report, verdicts = _match_column(merged, "greeting_standard")

        assert (report.n, report.tp, report.fp) == (4, 2, 2)
        # The verdicts must be the same comparison the counts came from, not a second one.
        assert verdicts == [MATCH_TRUE, MATCH_FALSE, MATCH_TRUE, MATCH_FALSE]
        assert verdicts.count(MATCH_TRUE) == report.tp

    def test_case_difference_is_a_match(self):
        gt = build_sheet(greeting_standard=["meet", "BELOW"])
        result = build_sheet(greeting_standard=["Meet", "Below"])
        merged, _, _ = _merge(gt, result)

        report, verdicts = _match_column(merged, "greeting_standard")

        assert report.tp == 2
        assert verdicts == [MATCH_TRUE, MATCH_TRUE]

    def test_value_outside_the_allowed_grades_is_a_mismatch_not_an_exclusion(self):
        gt = build_sheet(overall_sentiment=["Neutral", "a Thai narrative, not a label", "Positive"])
        result = build_sheet(overall_sentiment=["Neutral", "Neutral", "Positive"])
        merged, _, _ = _merge(gt, result)

        report, verdicts = _match_column(merged, "overall_sentiment")

        # The difference from google_confusion_matrix, and the reason N is constant: that row is
        # scored as a disagreement rather than dropped.
        assert report.n == 3
        assert report.tp == 2
        assert report.fp == 1
        assert verdicts == [MATCH_TRUE, MATCH_FALSE, MATCH_TRUE]

    def test_two_blanks_count_as_a_match_and_are_logged(self, capsys):
        gt = build_sheet(wrap_up=["", "", "Meet"])
        result = build_sheet(wrap_up=["", "", "Meet"])
        merged, _, _ = _merge(gt, result)

        report, verdicts = _match_column(merged, "wrap_up")

        # The documented consequence of excluding nothing. Pinned so it cannot change unnoticed,
        # and so the warning that makes it visible cannot be dropped as noise.
        assert report.accuracy == 1.0
        assert verdicts == [MATCH_TRUE, MATCH_TRUE, MATCH_TRUE]

        # capsys rather than caplog: structlog writes through its own logger factory, not the
        # stdlib logging caplog hooks, so a caplog assertion here passes vacuously.
        captured = capsys.readouterr()
        record = captured.err + captured.out
        assert "blank_both_sides" in record
        assert "wrap_up" in record
        # The count, never the value -- a stray cell in one of these columns is call content.
        assert "rows=2" in record or '"rows":2' in record

    def test_a_blank_on_one_side_only_is_a_mismatch(self):
        gt = build_sheet(wrap_up=["Meet", "Meet"])
        result = build_sheet(wrap_up=["", "Meet"])
        merged, _, _ = _merge(gt, result)

        report, _ = _match_column(merged, "wrap_up")

        assert report.tp == 1
        assert report.fp == 1


class TestEvaluateBlocks:
    def test_qa_block_carries_the_reference_column_set_in_order(self):
        blocks = evaluate(build_sheet(), build_sheet(), embedder=stub_embedder)

        assert list(blocks[BLOCK_QA].columns) == [
            "Column",
            "N",
            "TP",
            "FP",
            "FN",
            "TN",
            "Accuracy",
            "Precision",
            "Recall",
            "F1",
        ]
        assert list(blocks[BLOCK_SENTIMENT].columns) == list(blocks[BLOCK_QA].columns)

    def test_every_qa_and_sentiment_column_gets_a_row(self):
        blocks = evaluate(build_sheet(), build_sheet(), embedder=stub_embedder)

        assert list(blocks[BLOCK_QA]["Column"]) == list(QA_CRITERIA)
        assert list(blocks[BLOCK_SENTIMENT]["Column"]) == list(SENTIMENT_COLUMNS)

    def test_self_comparison_scores_one_everywhere(self):
        """The strongest check available: a sheet against itself must be perfect.

        A metric that cannot reach its own maximum on a perfect answer is broken, and this is the
        check that caught exactly that in google_confusion_matrix's macro-F1.
        """
        sheet = build_sheet(
            greeting_standard=["Meet", "Below", "N/A", "meet"],
            overall_sentiment=["Positive", "Negative", "Neutral", "Neutral"],
            call_type=["Enquiry", "Sale, Enquiry", "Retention", "Complaint"],
        )

        blocks = evaluate(sheet, sheet.copy(), embedder=stub_embedder)

        for title in (BLOCK_QA, BLOCK_SENTIMENT):
            block = blocks[title]
            assert list(block["TP"]) == list(block["N"])
            assert set(block["FP"]) == {0}
            for column in ("Accuracy", "Precision", "Recall", "F1"):
                assert set(block[column]) == {1.0}, f"{title} {column}"

        overall = blocks[BLOCK_OVERALL]
        for family in (FAMILY_QA, FAMILY_SENTIMENT, FAMILY_CALL_TYPE):
            row = frame_row(overall, "Family", family)
            assert row["Accuracy"] == 1.0
            assert row["F1"] == 1.0

    def test_n_is_the_matched_row_count_on_every_column(self):
        gt = build_sheet(
            greeting_standard=["Meet", "not a grade", "N/A", "Meet"],
            overall_sentiment=["", "Neutral", "Neutral", "Neutral"],
        )
        result = build_sheet(greeting_standard=["Meet", "Meet", "Meet", "Meet"])

        blocks = evaluate(gt, result, embedder=stub_embedder)

        # The guarantee the "nothing excluded" rule buys: no column can quietly score itself on a
        # smaller sample than the one beside it.
        assert set(blocks[BLOCK_QA]["N"]) == {4}
        assert set(blocks[BLOCK_SENTIMENT]["N"]) == {4}

    def test_family_accuracy_is_the_mean_of_its_columns(self):
        gt = build_sheet(
            greeting_standard=["Meet", "Below", "Meet", "Meet"],
            manners=["Meet", "Meet", "Below", "Below"],
        )
        result = build_sheet()

        blocks = evaluate(gt, result, embedder=stub_embedder)
        columns = blocks[BLOCK_QA]
        row = frame_row(blocks[BLOCK_OVERALL], "Family", FAMILY_QA)

        # Every column has the same N, so the mean of the accuracies is the pooled match rate
        # exactly -- 20 of 22 columns perfect, one at 0.75, one at 0.50.
        assert row["Accuracy"] == pytest.approx(round(columns["Accuracy"].mean(), 4))

    def test_overall_f1_column_is_named_f1_not_macro_f1(self):
        blocks = evaluate(build_sheet(), build_sheet(), embedder=stub_embedder)

        # It no longer means one thing down the column -- QA and Sentiment carry 2a/(1+a), Call
        # Type a genuine macro average -- so the heading must not claim otherwise.
        assert "F1" in blocks[BLOCK_OVERALL].columns
        assert "Macro-F1" not in blocks[BLOCK_OVERALL].columns

    def test_call_type_block_is_unchanged(self):
        gt = build_sheet(call_type=["Enquiry", "Sale", "Enquiry, Sale", "Retention"])
        result = build_sheet(call_type=["Enquiry", "Enquiry", "Sale, Enquiry", "Retention"])

        blocks = evaluate(gt, result, embedder=stub_embedder)
        metrics = dict(
            zip(
                blocks[BLOCK_CALL_TYPE_SUMMARY]["Metric"],
                blocks[BLOCK_CALL_TYPE_SUMMARY]["Value"],
                strict=True,
            )
        )

        # Still the multi-label report, not a match rate: the per-label detail is the whole reason
        # this one family was left alone.
        assert "Micro-F1" in metrics
        assert "Macro-F1 classes" in metrics
        assert metrics["Exact set match"] == 0.75

    def test_summary_row_stays_not_applicable(self):
        blocks = evaluate(build_sheet(), build_sheet(), embedder=stub_embedder)
        row = frame_row(blocks[BLOCK_OVERALL], "Family", "Summary Story")

        assert row["Accuracy"] == NOT_APPLICABLE
        assert row["F1"] == NOT_APPLICABLE

    def test_legend_explains_the_degenerate_columns(self):
        blocks = evaluate(build_sheet(), build_sheet(), embedder=stub_embedder)
        legend = dict(
            zip(blocks[BLOCK_LEGEND]["Term"], blocks[BLOCK_LEGEND]["Meaning"], strict=True)
        )

        # Printing Precision, Recall and F1 without saying what they are is the one way this
        # sheet could actively mislead, so the legend entries are part of the contract.
        assert "SAME NUMBER AS ACCURACY" in legend["Precision"]
        assert "1.0000" in legend["Recall"]
        assert "2 x Accuracy / (1 + Accuracy)" in legend["F1"]
        assert "Always 0" in legend["FN and TN"]


class TestCompareRows:
    def test_one_row_per_matched_call_in_join_order(self):
        gt = build_sheet(rows=3)
        result = build_sheet(rows=3)

        frames = compare_rows(gt, result)

        for key in COMPARE_KEYS:
            assert len(frames[key]) == 3
            assert list(frames[key]["No"]) == [1, 2, 3]
            assert list(frames[key][KEY_COLUMN]) == ["call_0.wav", "call_1.wav", "call_2.wav"]

    def test_qa_sheet_has_a_triplet_per_criterion(self):
        frames = compare_rows(build_sheet(), build_sheet())
        columns = list(frames[COMPARE_QA].columns)

        assert columns[:2] == ["No", KEY_COLUMN]
        assert len(columns) == 2 + len(QA_CRITERIA) * 3
        for criterion in QA_CRITERIA:
            assert [f"{criterion} GT", f"{criterion} AI", f"{criterion} Compare"] == columns[
                columns.index(f"{criterion} GT") : columns.index(f"{criterion} GT") + 3
            ]

    def test_cells_hold_the_raw_value_not_the_folded_one(self):
        gt = build_sheet(greeting_standard=["meet", "Below"])
        result = build_sheet(greeting_standard=["Meet", "Below"])

        frame = compare_rows(gt, result)[COMPARE_QA]

        # The sheet is evidence, so it must show what the workbook held. Folding it here would
        # make a reader chasing a verdict see a value that was never entered.
        assert list(frame["greeting_standard GT"]) == ["meet", "Below"]
        assert list(frame["greeting_standard AI"]) == ["Meet", "Below"]
        assert list(frame["greeting_standard Compare"]) == [MATCH_TRUE, MATCH_TRUE]

    def test_verdicts_reconcile_with_the_dashboard(self):
        gt = build_sheet(greeting_standard=["Meet", "Below", "N/A", "Meet"])
        result = build_sheet(greeting_standard=["Meet", "Meet", "N/A", "Below"])

        blocks = evaluate(gt, result, embedder=stub_embedder)
        frame = compare_rows(gt, result)[COMPARE_QA]
        row = frame_row(blocks[BLOCK_QA], "Column", "greeting_standard")

        # The compare sheet is the evidence for the block's counts; if they can disagree, it is
        # evidence for something else.
        assert list(frame["greeting_standard Compare"]).count(MATCH_TRUE) == row["TP"]
        assert list(frame["greeting_standard Compare"]).count(MATCH_FALSE) == row["FP"]

    def test_call_type_compares_sets_and_ignores_order(self):
        gt = build_sheet(call_type=["Enquiry, Sale", "Enquiry", "Sale"])
        result = build_sheet(call_type=["Sale,Enquiry", "Enquiry, Sale", "Sale"])

        frame = compare_rows(gt, result)[COMPARE_CALL_TYPE]

        assert list(frame["call_type Compare"]) == [MATCH_TRUE, MATCH_FALSE, MATCH_TRUE]
        # The raw text is still what is printed, so an unrecognised token stays visible.
        assert frame["call_type GT"][0] == "Enquiry, Sale"

    def test_empty_join_keeps_the_column_shape(self):
        gt = build_sheet(rows=2)
        result = build_sheet(rows=2)
        result[KEY_COLUMN] = ["other_a.wav", "other_b.wav"]

        frames = compare_rows(gt, result)

        # Without this the sheet would lose its header and read as a successful empty run.
        assert len(frames[COMPARE_QA]) == 0
        assert len(frames[COMPARE_QA].columns) == 2 + len(QA_CRITERIA) * 3
        assert len(frames[COMPARE_SENTIMENT].columns) == 2 + len(SENTIMENT_COLUMNS) * 3


class TestResolveSheetNames:
    BASES = (
        "Evaluation Dashboard",
        "QA Compare Result",
        "Sentiment Compare Result",
        "Call Type Compare Result",
    )

    def test_returns_the_bases_when_all_are_free(self):
        assert resolve_sheet_names(["Sheet1"], self.BASES) == list(self.BASES)

    def test_one_taken_name_moves_all_four(self):
        resolved = resolve_sheet_names(["QA Compare Result"], self.BASES)

        # The four sheets are only readable as a set. A dashboard paired with a verdict sheet
        # from a different run is worse than an ugly suffix.
        assert resolved == [f"{base}_1" for base in self.BASES]

    def test_skips_a_suffix_any_base_has_taken(self):
        existing = ["Evaluation Dashboard", "Sentiment Compare Result_1", "QA Compare Result_2"]

        resolved = resolve_sheet_names(existing, self.BASES)

        assert resolved == [f"{base}_3" for base in self.BASES]

    def test_config_supplies_all_four_names(self):
        config = Config()

        assert [
            config.eval_sheet_name,
            config.qa_compare_sheet_name,
            config.sentiment_compare_sheet_name,
            config.call_type_compare_sheet_name,
        ] == list(self.BASES)


class TestColumnFormatting:
    def test_n_is_a_count_not_a_ratio(self):
        # Left out of COUNT_COLUMNS it gets no format at all, and a stray ratio match would paint
        # 152 solid green on a 0-1 colour scale.
        assert _is_count_column("N")
        assert not _is_ratio_column("N")

    @pytest.mark.parametrize("name", ["Accuracy", "Precision", "Recall", "F1"])
    def test_the_four_scores_are_ratios(self, name):
        assert _is_ratio_column(name)
        assert not _is_count_column(name)

    @pytest.mark.parametrize("name", ["TP", "FP", "FN", "TN"])
    def test_the_four_counts_are_counts(self, name):
        assert _is_count_column(name)
        assert not _is_ratio_column(name)


@pytest.fixture(scope="module")
def rendered():
    """One workbook carrying all four appended sheets, built once for the rendering tests."""
    gt = build_sheet(greeting_standard=["Meet", "Below", "N/A", "meet"])
    result = build_sheet(greeting_standard=["Meet", "Meet", "N/A", "Meet"])
    original = starting_workbook(gt, result)

    blocks = evaluate(gt, result, embedder=stub_embedder)
    frames = compare_rows(gt, result)
    names = resolve_sheet_names(
        load_workbook(io.BytesIO(original)).sheetnames, TestResolveSheetNames.BASES
    )
    updated = write_dashboard(original, blocks, names[0])
    updated = write_comparison_sheets(
        updated, frames, dict(zip(COMPARE_KEYS, names[1:], strict=True))
    )
    return original, updated, names


class TestRenderedWorkbook:
    def test_all_four_sheets_are_appended(self, rendered):
        _, updated, names = rendered

        book = load_workbook(io.BytesIO(updated))

        assert book.sheetnames == [GT_SHEET, RESULT_SHEET, *names]

    def test_source_sheets_survive_untouched(self, rendered):
        original, updated, _ = rendered

        before = pd.read_excel(io.BytesIO(original), sheet_name=GT_SHEET)
        after = pd.read_excel(io.BytesIO(updated), sheet_name=GT_SHEET)

        assert before.equals(after)

    def test_compare_sheet_has_a_two_row_merged_header(self, rendered):
        _, updated, names = rendered
        sheet = load_workbook(io.BytesIO(updated))[names[1]]

        assert sheet.cell(row=1, column=1).value == "No"
        assert sheet.cell(row=1, column=2).value == KEY_COLUMN
        assert sheet.cell(row=1, column=3).value == QA_CRITERIA[0]
        # Row 1 names the group, row 2 names the three cells under it.
        assert [sheet.cell(row=2, column=c).value for c in (3, 4, 5)] == ["GT", "AI", "Compare"]
        assert "C1:E1" in {str(r) for r in sheet.merged_cells.ranges}
        assert "A1:A2" in {str(r) for r in sheet.merged_cells.ranges}

    def test_compare_sheet_freezes_the_header_and_the_key(self, rendered):
        _, updated, names = rendered
        sheet = load_workbook(io.BytesIO(updated))[names[1]]

        # Both header rows and both index columns: with 22 triplets the criterion name is off
        # screen long before its Compare cell is.
        assert sheet.freeze_panes == "C3"
        assert sheet.auto_filter.ref.startswith("A2:")

    def test_mismatch_cells_are_filled_red_and_matches_green(self, rendered):
        _, updated, names = rendered
        sheet = load_workbook(io.BytesIO(updated))[names[1]]

        verdicts = {sheet.cell(row=r, column=5).value: r for r in range(3, 7)}
        matched = sheet.cell(row=verdicts[MATCH_TRUE], column=5)
        missed = sheet.cell(row=verdicts[MATCH_FALSE], column=5)

        assert missed.fill.fgColor.rgb.endswith("FFC7CE")
        assert matched.fill.fgColor.rgb.endswith("C6EFCE")

    def test_dashboard_formats_n_as_an_integer_and_accuracy_at_four_places(self, rendered):
        _, updated, names = rendered
        sheet = load_workbook(io.BytesIO(updated))[names[0]]

        header = next(row for row in sheet.iter_rows() if row[0].value == "Column")
        first_data = header[0].row + 1

        assert [cell.value for cell in header[:6]] == ["Column", "N", "TP", "FP", "FN", "TN"]
        assert sheet.cell(row=first_data, column=2).number_format == "0"
        assert sheet.cell(row=first_data, column=7).number_format == "0.0000"

    def test_writing_twice_produces_a_second_suffixed_set(self, rendered):
        _, updated, _ = rendered
        gt = build_sheet(rows=2)
        result = build_sheet(rows=2)

        second = resolve_sheet_names(
            load_workbook(io.BytesIO(updated)).sheetnames, TestResolveSheetNames.BASES
        )
        again = write_dashboard(
            updated, evaluate(gt, result, embedder=stub_embedder), second[0]
        )
        again = write_comparison_sheets(
            again, compare_rows(gt, result), dict(zip(COMPARE_KEYS, second[1:], strict=True))
        )

        assert second == [f"{base}_1" for base in TestResolveSheetNames.BASES]
        assert load_workbook(io.BytesIO(again)).sheetnames[-4:] == second

    def test_writing_over_an_existing_compare_sheet_raises(self, rendered):
        _, updated, names = rendered
        gt = build_sheet(rows=2)

        with pytest.raises(ValueError, match="already exists"):
            write_comparison_sheets(
                updated,
                compare_rows(gt, gt.copy()),
                dict(zip(COMPARE_KEYS, names[1:], strict=True)),
            )
