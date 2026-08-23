"""Tests for the documents confusion-matrix evaluation layer.

Synthetic sheets throughout, built by :func:`build_sheet`, plus a skip-guarded class over the
labelled run artifact under ``debug/``.

What is worth pinning here is the handful of decisions that separate this module from the
exact-match sibling: rows outside the class vocabulary are excluded and counted rather than
scored as mismatches, the vocabularies fold case (and the flags fold ``1``/``0``), and a class
neither side used reads ``n/a`` rather than a red 0.0000.
"""

import io
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from src.google_model.documents.google_confusion_matrix import (
    BLOCK_DOC_TYPE,
    BLOCK_DOC_TYPE_LABELS,
    BLOCK_FLAGS,
    BLOCK_HEADER,
    BLOCK_LEGEND,
    BOOLEAN_COLUMNS,
    BOOLEAN_LABELS,
    DOC_TYPE_COLUMN,
    DOC_TYPE_LABELS,
    KEY_COLUMN,
    NOT_APPLICABLE,
    _canonical_bool,
    _canonical_doc_type,
    _merge,
    _score_single_label,
    evaluate,
    resolve_sheet_name,
    write_dashboard,
)

GT_SHEET = "Doc - Groundtruth"
RESULT_SHEET = "Doc - Google Result"

WORKBOOK = Path(__file__).resolve().parents[1] / "debug" / "model_comparison_2026-08-09_18-25-52_gt.xlsx"
READ_KWARGS = {"dtype": str, "header": 0, "keep_default_na": False, "na_values": [""]}


def build_sheet(rows=4, **columns):
    """A minimal sheet carrying every column evaluate() reads.

    Defaults agree on every column, so a test supplies only the column it is about and any
    disagreement in the result is the one it introduced. Column headers carry spaces, so
    overrides are passed ``**{"Document type": [...]}``.
    """
    if columns:
        rows = len(next(iter(columns.values())))

    frame = {
        KEY_COLUMN: [f"doc_{i}_p1.pdf#1" for i in range(rows)],
        DOC_TYPE_COLUMN: ["Receipt"] * rows,
    }
    frame.update({column: ["True"] * rows for column in BOOLEAN_COLUMNS})
    frame.update(columns)
    return pd.DataFrame(frame)


def frame_row(frame, key_column, value):
    """The single row of a block whose key column holds ``value``."""
    matches = frame[frame[key_column] == value]
    assert len(matches) == 1, f"{value!r} matched {len(matches)} rows"
    return matches.iloc[0]


class TestCanonicalise:
    def test_doc_type_folds_case_only(self):
        assert _canonical_doc_type("taxinvoice") == "TaxInvoice"
        assert _canonical_doc_type(" RECEIPT ") == "Receipt"
        # No synonym map, on purpose: folding "Invoice" into "TaxInvoice" would be a business
        # judgement, and metric code is the wrong place to make one.
        assert _canonical_doc_type("Invoice") is None

    def test_blank_and_unknown_doc_types_are_none(self):
        assert _canonical_doc_type("") is None
        assert _canonical_doc_type(None) is None
        assert _canonical_doc_type("a Thai narrative, not a class") is None

    @pytest.mark.parametrize(
        "value,expected",
        [("True", "True"), ("TRUE", "True"), ("1", "True"), ("false", "False"), ("0", "False")],
    )
    def test_bool_vocabulary(self, value, expected):
        assert _canonical_bool(value) == expected

    def test_blank_and_unknown_bools_are_none(self):
        assert _canonical_bool("") is None
        assert _canonical_bool("yes") is None


class TestScoreSingleLabel:
    def test_out_of_vocabulary_rows_are_dropped_and_counted(self, capsys):
        gt = build_sheet(
            **{DOC_TYPE_COLUMN: ["Receipt", "not a class", "Other", ""]}
        )
        result = build_sheet(**{DOC_TYPE_COLUMN: ["Receipt", "Receipt", "Other", "Receipt"]})
        merged, _, _ = _merge(gt, result)

        report, dropped = _score_single_label(
            merged, DOC_TYPE_COLUMN, DOC_TYPE_LABELS, _canonical_doc_type
        )

        # The deliberate difference from the exact-match module, which scores those two rows as
        # mismatches: here they measure a data-entry fault, not the model.
        assert (report.n, dropped) == (2, 2)
        assert report.accuracy == 1.0
        # The full label set survives the filtering -- the report is over the vocabulary, not
        # over whatever values happened to appear.
        assert list(report.labels) == list(DOC_TYPE_LABELS)

        # capsys rather than caplog: structlog writes through its own logger factory, not the
        # stdlib logging caplog hooks, so a caplog assertion here passes vacuously.
        captured = capsys.readouterr()
        record = captured.err + captured.out
        assert "column.excluded" in record
        assert "not a class" not in record  # the value itself never reaches the log

    def test_case_difference_is_scored_not_dropped(self):
        gt = build_sheet(**{DOC_TYPE_COLUMN: ["receipt", "TAXINVOICE"]})
        result = build_sheet(**{DOC_TYPE_COLUMN: ["Receipt", "TaxInvoice"]})
        merged, _, _ = _merge(gt, result)

        report, dropped = _score_single_label(
            merged, DOC_TYPE_COLUMN, DOC_TYPE_LABELS, _canonical_doc_type
        )

        assert (report.n, dropped) == (2, 0)
        assert report.accuracy == 1.0

    def test_flags_score_one_zero_against_true_false(self):
        gt = build_sheet(**{"Stamp": ["1", "0", "True"]})
        result = build_sheet(**{"Stamp": ["True", "False", "True"]})
        merged, _, _ = _merge(gt, result)

        report, dropped = _score_single_label(merged, "Stamp", BOOLEAN_LABELS, _canonical_bool)

        assert (report.n, dropped) == (3, 0)
        assert report.accuracy == 1.0


class TestEvaluate:
    def test_blocks_arrive_in_sheet_order(self):
        blocks = evaluate(build_sheet(), build_sheet())

        assert list(blocks) == [
            BLOCK_HEADER,
            BLOCK_DOC_TYPE,
            BLOCK_DOC_TYPE_LABELS,
            BLOCK_FLAGS,
            BLOCK_LEGEND,
        ]

    def test_self_comparison_scores_one(self):
        """The strongest check available: a sheet against itself must be perfect."""
        sheet = build_sheet(
            **{
                DOC_TYPE_COLUMN: ["Receipt", "TaxInvoice", "Suspicious", "Receipt"],
                "Copy": ["True", "False", "True", "False"],
            }
        )

        blocks = evaluate(sheet, sheet.copy())
        metrics = dict(
            zip(blocks[BLOCK_DOC_TYPE]["Metric"], blocks[BLOCK_DOC_TYPE]["Value"], strict=True)
        )

        assert metrics["Accuracy"] == 1.0
        assert metrics["Macro-F1"] == 1.0
        assert metrics["Scored"] == 4
        assert metrics["Excluded (blank or unknown)"] == 0

        flags = blocks[BLOCK_FLAGS]
        assert set(flags["Accuracy"]) == {1.0}

    def test_per_label_block_carries_all_six_classes(self):
        blocks = evaluate(build_sheet(), build_sheet())
        labels_block = blocks[BLOCK_DOC_TYPE_LABELS]

        assert list(labels_block["Label"]) == list(DOC_TYPE_LABELS)

    def test_a_class_nobody_used_reads_not_applicable(self):
        # Only "Receipt" appears; IDCard scores 0.0000 on every metric purely because every
        # denominator was zero, and printing those zeros would paint a red cell on a column
        # that is in fact perfect.
        blocks = evaluate(build_sheet(), build_sheet())
        row = frame_row(blocks[BLOCK_DOC_TYPE_LABELS], "Label", "IDCard")

        assert row["Precision"] == NOT_APPLICABLE
        assert row["F1"] == NOT_APPLICABLE
        assert row["Support"] == 0

    def test_a_class_only_the_model_invented_stays_in_at_zero(self):
        gt = build_sheet(**{DOC_TYPE_COLUMN: ["Receipt", "Receipt", "Receipt"]})
        result = build_sheet(**{DOC_TYPE_COLUMN: ["Receipt", "Quotation", "Receipt"]})

        blocks = evaluate(gt, result)
        row = frame_row(blocks[BLOCK_DOC_TYPE_LABELS], "Label", "Quotation")

        # Inventing a class the human never gave is a real error, so it is not blanked away.
        assert row["Support"] == 0
        assert row["FP"] == 1
        assert row["F1"] == 0.0

    def test_excluded_rows_reach_the_summary_block(self):
        gt = build_sheet(**{DOC_TYPE_COLUMN: ["Receipt", "not a class", "Receipt"]})
        result = build_sheet(rows=3)

        blocks = evaluate(gt, result)
        metrics = dict(
            zip(blocks[BLOCK_DOC_TYPE]["Metric"], blocks[BLOCK_DOC_TYPE]["Value"], strict=True)
        )

        assert metrics["Scored"] == 2
        assert metrics["Excluded (blank or unknown)"] == 1

    def test_flags_block_carries_per_class_cells_for_both_labels(self):
        blocks = evaluate(build_sheet(), build_sheet())
        flags = blocks[BLOCK_FLAGS]

        assert list(flags["Column"]) == list(BOOLEAN_COLUMNS)
        for label in BOOLEAN_LABELS:
            for metric in ("TP", "TN", "FP", "FN", "Precision", "Recall", "F1", "Support"):
                assert f"{label} {metric}" in flags.columns

    def test_legend_warns_about_the_majority_class(self):
        blocks = evaluate(build_sheet(), build_sheet())
        legend = dict(
            zip(blocks[BLOCK_LEGEND]["Term"], blocks[BLOCK_LEGEND]["Meaning"], strict=True)
        )

        assert "NEVER READ IT WITHOUT MACRO-F1" in legend["Accuracy"]
        assert "n/a" in legend["Support 0"]


class TestResolveSheetName:
    def test_returns_the_base_when_free(self):
        assert resolve_sheet_name(["Sheet1"], "Evaluation Dashboard") == "Evaluation Dashboard"

    def test_suffixes_on_collision(self):
        existing = ["Evaluation Dashboard", "Evaluation Dashboard_1"]

        assert resolve_sheet_name(existing, "Evaluation Dashboard") == "Evaluation Dashboard_2"


@pytest.fixture(scope="module")
def rendered():
    """One workbook carrying the appended dashboard, built once for the rendering tests."""
    gt = build_sheet(**{DOC_TYPE_COLUMN: ["Receipt", "TaxInvoice", "Receipt", "Receipt"]})
    result = build_sheet(**{DOC_TYPE_COLUMN: ["Receipt", "Receipt", "Receipt", "Receipt"]})

    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        gt.to_excel(writer, sheet_name=GT_SHEET, index=False)
        result.to_excel(writer, sheet_name=RESULT_SHEET, index=False)
    original = buffer.getvalue()

    blocks = evaluate(gt, result)
    name = resolve_sheet_name(load_workbook(io.BytesIO(original)).sheetnames, "Evaluation Dashboard")
    updated = write_dashboard(original, blocks, name)
    return original, updated, name


class TestRenderedWorkbook:
    def test_the_dashboard_is_appended(self, rendered):
        _, updated, name = rendered

        assert load_workbook(io.BytesIO(updated)).sheetnames == [GT_SHEET, RESULT_SHEET, name]

    def test_source_sheets_survive_untouched(self, rendered):
        original, updated, _ = rendered

        before = pd.read_excel(io.BytesIO(original), sheet_name=GT_SHEET)
        after = pd.read_excel(io.BytesIO(updated), sheet_name=GT_SHEET)

        assert before.equals(after)

    def test_metric_value_rows_format_by_metric_name(self, rendered):
        _, updated, name = rendered
        sheet = load_workbook(io.BytesIO(updated))[name]

        # First occurrence wins: the legend also carries an "Accuracy" row in column A, and it
        # is deliberately unformatted -- the Metric/Value block sits above it on the sheet.
        rows = {}
        for row in sheet.iter_rows():
            if row[0].value:
                rows.setdefault(row[0].value, row)

        # A Metric/Value block mixes ratios with counts in one column, so the format is decided
        # per row: Accuracy at four places, Scored as a bare integer.
        assert rows["Accuracy"][1].number_format == "0.0000"
        assert rows["Scored"][1].number_format == "0"

    def test_not_applicable_cells_are_muted_not_scored(self, rendered):
        _, updated, name = rendered
        sheet = load_workbook(io.BytesIO(updated))[name]

        muted = [
            cell
            for row in sheet.iter_rows()
            for cell in row
            if cell.value == NOT_APPLICABLE and cell.font.italic
        ]

        # The unused classes' score cells must render greyed, or a blank reads as a zero.
        assert muted

    def test_writing_over_an_existing_sheet_raises(self, rendered):
        _, updated, name = rendered

        with pytest.raises(ValueError, match="already exists"):
            write_dashboard(updated, evaluate(build_sheet(), build_sheet()), name)


@pytest.fixture(scope="module")
def real_sheets():
    """The labelled run artifact's two sheets, when the file is present."""
    if not WORKBOOK.exists():
        pytest.skip(f"sample workbook not present: {WORKBOOK}")
    book = pd.ExcelFile(WORKBOOK, engine="openpyxl")
    gt = book.parse(GT_SHEET, **READ_KWARGS)
    result = book.parse(RESULT_SHEET, **READ_KWARGS)
    return gt, result


class TestRealWorkbook:
    def test_evaluate_runs_offline_over_the_real_frames(self, real_sheets):
        gt, result = real_sheets
        blocks = evaluate(gt, result)

        assert len(blocks[BLOCK_DOC_TYPE_LABELS]) == len(DOC_TYPE_LABELS)
        assert len(blocks[BLOCK_FLAGS]) == len(BOOLEAN_COLUMNS)

    def test_scored_plus_excluded_accounts_for_every_matched_row(self, real_sheets):
        gt, result = real_sheets
        merged, _, _ = _merge(gt, result)

        blocks = evaluate(gt, result)
        metrics = dict(
            zip(blocks[BLOCK_DOC_TYPE]["Metric"], blocks[BLOCK_DOC_TYPE]["Value"], strict=True)
        )

        assert metrics["Scored"] + metrics["Excluded (blank or unknown)"] == len(merged)
